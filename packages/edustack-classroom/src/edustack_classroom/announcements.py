import os
import re
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from datetime import timedelta
from pathlib import Path
from edustack_schedule.term_resolver import resolve_term

def generate_announcements(
    excel_file: str | Path,
    output_dir: str | Path = "Announcements",
    semester: str | None = None,
    year: int | None = None,
):
    excel_file = Path(excel_file)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if not semester or not year:
        semester, year = resolve_term()

    term_label = f"{semester} {year}"

    # Load data from Excel
    deliverables = pd.read_excel(excel_file, sheet_name="Deliverables")
    course_details = pd.read_excel(excel_file, sheet_name="Course Details")
    calendar = pd.read_excel(excel_file, sheet_name="Calendar")

    # Filter calendar for the current semester
    current_calendar = calendar[calendar["Semester"].str.contains(term_label, case=False, na=False)]
    if current_calendar.empty:
        print(f"Warning: No calendar entries found for term '{term_label}'. Falling back to all rows.")
        current_calendar = calendar

    # Helper function to set font for paragraphs or headings
    def set_font(paragraph, font_name="Cambria"):
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(11)

    # Helper function to format table cells
    def format_table_cell(cell, font_name="Cambria", font_size=11):
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = font_name
                run.font.size = Pt(font_size)

    # Extract week columns (columns starting with W followed by digits, like W01D6 or W01)
    week_columns = [col for col in deliverables.columns if re.match(r"^W\d+", col)]

    for week_col in week_columns:
        week_number = week_col[1:3]  # Two-digit week number (e.g., 01, 02)
        week_int = int(week_number)

        # Get the title and description for the current week's content
        try:
            lecture_title = course_details.loc[week_int - 1, "Title"]
            lecture_description = course_details.loc[week_int - 1, "Description"]
        except Exception:
            lecture_title = f"Lecture {week_int}"
            lecture_description = ""

        # Get start date columns per section
        start_date_column = f"W{week_number}D1"
        if start_date_column not in calendar.columns:
            print(f"Warning: Start date column {start_date_column} not found in Calendar sheet. Skipping week {week_number}.")
            continue

        # Extract sections and their open/close dates
        section_dates = []
        for _, row in current_calendar.iterrows():
            sec = row.get("Section")
            dt_val = row.get(start_date_column)
            if pd.notna(sec) and pd.notna(dt_val):
                open_date = pd.to_datetime(dt_val)
                close_date = open_date + timedelta(days=6)
                section_dates.append((sec, open_date, close_date))

        if not section_dates:
            print(f"Warning: No start dates found for week {week_number} in calendar. Skipping.")
            continue

        # Create Word Document
        document = Document()

        # Add the announcement heading
        h1_heading = document.add_heading(f"Announcements for Week {week_number}", level=1)
        h1_heading.runs[0].font.color.rgb = RGBColor(204, 0, 0)  # Red
        set_font(h1_heading, "Cambria")

        # Add general introduction
        intro = document.add_paragraph("Hello, everybody!")
        set_font(intro, "Cambria")
        intro.add_run(f"\nHere are the essential deliverables for week {week_int} of AD 688 {term_label}:").font.size = Pt(12)

        # Add Week Content Review section
        content_heading = document.add_heading(f"1. Week {week_int} Content Review", level=2)
        content_heading.runs[0].font.color.rgb = RGBColor(204, 0, 0)
        set_font(content_heading, "Cambria")

        content_paragraph = document.add_paragraph()
        title_run = content_paragraph.add_run(lecture_title)
        title_run.bold = True
        set_font(content_paragraph, "Cambria")
        content_paragraph.add_run(f"\n{lecture_description}").font.size = Pt(12)

        # Add In-Class Labs section
        labs_heading = document.add_heading("2. In-Class Labs", level=2)
        labs_heading.runs[0].font.color.rgb = RGBColor(204, 0, 0)
        set_font(labs_heading, "Cambria")

        # Add table for labs
        table = document.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        headers = ["Section", "Lab", "Open Date", "Close Date"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            format_table_cell(cell)

        # Populate table
        for sec, open_date, close_date in section_dates:
            row = table.add_row()
            row.cells[0].text = str(sec)
            row.cells[1].text = f"Lab {week_number}"
            row.cells[2].text = open_date.strftime('%A, %B %d, %Y')
            row.cells[3].text = close_date.strftime('%A, %B %d, %Y')
            for cell in row.cells:
                format_table_cell(cell)

        # Add Academy Labs Consultation section
        consultation_heading = document.add_heading(f"3. Academy Labs Consultation (Friday)", level=2)
        consultation_heading.runs[0].font.color.rgb = RGBColor(204, 0, 0)
        set_font(consultation_heading, "Cambria")

        consultation_paragraph = document.add_paragraph()
        consultation_paragraph.add_run("Zoom Consultation Room. Facilitators: Course Instructors and TAs.").bold = True
        set_font(consultation_paragraph, "Cambria")
        document.add_paragraph("Time: 9:00 AM - 11:00 AM (Friday).", style='List Bullet')
        document.add_paragraph("This session is an opportunity to ask questions, discuss assignments, and seek clarification on course materials.", style='List Bullet')
        document.add_paragraph("Zoom link will be available in the course materials section.", style='List Bullet')

        # Save document
        output_file = output_dir / f"Week_{week_number}_Announcement.docx"
        document.save(output_file)
        print(f"Generated {output_file}")
